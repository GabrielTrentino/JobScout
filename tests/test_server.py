import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document

from jobscout.server import serve_in_background


class ServerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.httpd = serve_in_background()
        host, port = cls.httpd.server_address
        cls.base = f"http://{host}:{port}"

    @classmethod
    def tearDownClass(cls) -> None:
        cls.httpd.shutdown()
        cls.httpd.server_close()

    def test_home_is_the_intake_screen(self) -> None:
        with urlopen(self.base + "/") as response:
            page = response.read().decode("utf-8")
        self.assertIn("Atividades do currículo", page)
        self.assertIn("Solte o PDF ou o DOCX aqui.", page)
        self.assertIn("/api/analisar", page)

    def test_post_docx_lists_activities(self) -> None:
        with TemporaryDirectory() as folder:
            path = Path(folder) / "cv.docx"
            document = Document()
            document.add_paragraph("Experiência profissional")
            document.add_paragraph("Analista | Oficina")
            document.add_paragraph("• Modelei o fluxo de pedidos do comercial")
            document.save(path)
            payload = path.read_bytes()
        body, content_type = _multipart("cv.docx", payload)
        request = Request(
            self.base + "/api/analisar",
            data=body,
            headers={"Content-Type": content_type},
            method="POST",
        )
        with urlopen(request) as response:
            result = json.loads(response.read().decode("utf-8"))
        self.assertEqual(result["modo"], "local")
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["ferramentas"], [])
        self.assertEqual(result["grupos"][0]["atividades"][0], "Modelei o fluxo de pedidos do comercial")

    def test_rejects_other_extensions(self) -> None:
        body, content_type = _multipart("notas.txt", b"texto")
        request = Request(
            self.base + "/api/analisar",
            data=body,
            headers={"Content-Type": content_type},
            method="POST",
        )
        with self.assertRaises(HTTPError) as caught:
            urlopen(request)
        self.assertEqual(caught.exception.code, 422)
        caught.exception.close()


def _multipart(filename: str, payload: bytes) -> tuple[bytes, str]:
    boundary = "----jobscouttest"
    head = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="arquivo"; filename="{filename}"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n"
    ).encode("utf-8")
    tail = f"\r\n--{boundary}--\r\n".encode("utf-8")
    return head + payload + tail, f"multipart/form-data; boundary={boundary}"


if __name__ == "__main__":
    unittest.main()
