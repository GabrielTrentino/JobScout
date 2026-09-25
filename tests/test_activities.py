import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from jobscout.activities import ResumeReadError, list_activities, list_tools, read_resume


CURRICULO = """
Ana Souza
Desenvolvedora backend
São Paulo, SP
ana@email.com

Resumo
Desenvolvedora com experiência em APIs e filas.

Experiência profissional

Desenvolvedora backend | Acme
São Paulo — Brasil
jan 2022 — atual
• Desenvolvi APIs REST em FastAPI para o faturamento
• Implementei filas com Redis e reduzi o tempo de processamento
Liderei a migração do banco para PostgreSQL

Analista de sistemas — Beta
2019 — 2021
- Modelei o fluxo de pedidos do comercial
- Automatizei relatórios mensais em Python

Projetos
Painel de operações | Pessoal
• Construí um painel interno para acompanhar filas

Formação
Bacharelado em Ciência da Computação — USP — 2018

Habilidades
Python, FastAPI, PostgreSQL, Docker
"""


class ActivitiesTest(unittest.TestCase):
    def test_lists_roles_and_skips_summary_and_skills(self) -> None:
        groups = list_activities(CURRICULO)
        contexts = [group["contexto"] for group in groups]
        self.assertEqual(
            contexts,
            [
                "Desenvolvedora backend — Acme",
                "Analista de sistemas — Beta",
                "Painel de operações — Pessoal",
            ],
        )
        self.assertEqual(groups[0]["periodo"], "jan 2022 — atual")
        self.assertEqual(groups[1]["periodo"], "2019 — 2021")
        self.assertEqual(groups[0]["secao"], "Experiência")
        self.assertEqual(
            groups[0]["atividades"],
            [
                "Desenvolvi APIs REST em FastAPI para o faturamento",
                "Implementei filas com Redis e reduzi o tempo de processamento",
                "Liderei a migração do banco para PostgreSQL",
            ],
        )
        self.assertEqual(groups[2]["secao"], "Projetos")
        flat = " ".join(item for group in groups for item in group["atividades"])
        self.assertNotIn("Bacharelado", flat)
        self.assertNotIn("experiência em APIs", flat)
        self.assertNotIn("PostgreSQL, Docker", flat)

    def test_fallback_without_section_headers(self) -> None:
        text = """
        Marina Lima
        marina@email.com
        (11) 98888-7777
        • Desenvolvi serviços de pagamento em Python
        • Entreguei a integração com o parceiro bancário
        """
        groups = list_activities(text)
        activities = groups[0]["atividades"]
        self.assertEqual(len(activities), 2)
        self.assertEqual(groups[0]["contexto"], "")
        self.assertNotIn("marina@email.com", " ".join(activities))

    def test_lists_tools_and_keeps_long_roles_apart(self) -> None:
        text = """
        Stack técnica
        Linguagens Python (7 anos) · PySpark (4 anos) · SQL · PL/SQL · Linux
        Cloud GCP — Airflow · BigQuery · Cloud Functions · Cloud Scheduler · Azure (2 anos) · AWS CCP
        Dados Power BI · Informatica Cloud (IICS) · Databricks · dbt · Apache Spark
        Dev Git/GitHub · Bitbucket · Jira

        Experiência profissional
        Engenheiro de Dados · Planal Lubrificantes LTDA · Mar/2024 – Atual · Goiânia, GO
        • Estruturei do zero o ecossistema de dados com pipelines e dashboards financeiros.
        • Automatizei cálculo de pedidos via interface web eliminando a precificação manual.

        Engenheiro de Dados (Trainee → Pleno) · Accenture & Tenbu · Jul/2021 – Mar/2024 · Home Office
        Analytics Engineer em DataLakes de missão crítica para seguros e saúde (GCP e Azure).
        • Construí e sustentei 40+ DAGs em Apache Airflow (GCP Composer) para pipelines.
        • Projetei DataLake premiado em governança com arquitetura Raw para Refined.

        Analista de Dados Jr. · Planal Lubrificantes LTDA · Jan/2019 – Jun/2021 · Goiânia, GO
        • Automatizei extração de dados de estoque e vendas com Python e PL/SQL.
        """
        tools = list_tools(text)
        self.assertEqual(
            tools,
            [
                "Python",
                "PySpark",
                "SQL",
                "PL/SQL",
                "Linux",
                "GCP",
                "Apache Airflow",
                "BigQuery",
                "Cloud Functions",
                "Cloud Scheduler",
                "Azure",
                "AWS",
                "Power BI",
                "Informatica Cloud",
                "Databricks",
                "dbt",
                "Apache Spark",
                "Git",
                "GitHub",
                "Bitbucket",
                "Jira",
            ],
        )
        groups = list_activities(text)
        contexts = [group["contexto"] for group in groups]
        self.assertEqual(len(contexts), 3)
        self.assertIn("Planal", contexts[0])
        self.assertIn("Accenture", contexts[1])
        self.assertIn("Analista", contexts[2])
        accenture = " ".join(groups[1]["atividades"])
        self.assertIn("Apache Airflow", accenture)
        self.assertNotIn("Apache Airflow", " ".join(groups[0]["atividades"]))

    def test_empty_text_has_no_activities(self) -> None:
        self.assertEqual(list_activities("   "), [])

    def test_docx_roundtrip(self) -> None:
        document = Document()
        document.add_paragraph("Experiência profissional")
        document.add_paragraph("Engenheira de software | Nuvem")
        document.add_paragraph("2021 — 2024")
        document.add_paragraph(
            "Implementei o fluxo de cadastro dos clientes",
            style="List Bullet",
        )
        document.add_paragraph("Formação")
        document.add_paragraph("Engenharia de computação")
        with TemporaryDirectory() as folder:
            path = Path(folder) / "curriculo.docx"
            document.save(path)
            groups = list_activities(read_resume(path))
        self.assertEqual(groups[0]["contexto"], "Engenheira de software — Nuvem")
        self.assertEqual(
            groups[0]["atividades"],
            ["Implementei o fluxo de cadastro dos clientes"],
        )

    def test_rejects_plain_text_file(self) -> None:
        with TemporaryDirectory() as folder:
            path = Path(folder) / "curriculo.txt"
            path.write_text("Experiência\n• Desenvolvi APIs", encoding="utf-8")
            with self.assertRaises(ResumeReadError):
                read_resume(path)

    def test_pdf_text(self) -> None:
        from jobscout.activities import _pdf_bytes

        payload = _simple_pdf(
            [
                "EXPERIENCIA PROFISSIONAL",
                "Developer | North",
                "2020 - 2023",
                "- Developed billing APIs for the finance team",
                "Habilidades",
                "Python, SQL",
            ]
        )
        groups = list_activities(_pdf_bytes(payload))
        self.assertEqual(groups[0]["atividades"], ["Developed billing APIs for the finance team"])


def _simple_pdf(lines: list[str]) -> bytes:
    commands = ["BT", "/F1 11 Tf", "50 760 Td", "16 TL"]
    for index, line in enumerate(lines):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if index == 0:
            commands.append(f"({escaped}) Tj")
        else:
            commands.append(f"T* ({escaped}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
        f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    objects[3] = b"4 0 obj\n" + objects[3]
    buffer = BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets = [0]
    for item in objects:
        offsets.append(buffer.tell())
        buffer.write(item)
    xref = buffer.tell()
    buffer.write(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    buffer.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        buffer.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    buffer.write(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
