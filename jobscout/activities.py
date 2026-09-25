"""Extrai atividades de um currículo sem chamar nenhum modelo."""

from __future__ import annotations

import re
import unicodedata
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 8 * 1024 * 1024

SECTION_STARTS = (
    ("experiencia", re.compile(
        r"^(experi[eê]ncias?(?:\s+profission(?:al|ais))?|hist[oó]rico\s+profissional|"
        r"atua[cç][aã]o\s+profissional|work\s+experience|professional\s+experience|"
        r"employment(?:\s+history)?)\s*:?\s*$",
        re.IGNORECASE,
    )),
    ("projetos", re.compile(
        r"^(projetos?(?:\s+relevantes)?|projects?)\s*:?\s*$",
        re.IGNORECASE,
    )),
)

SECTION_STOP = re.compile(
    r"^(forma[cç][aã]o(?:\s+acad[eê]mica)?|educa[cç][aã]o|education|academic(?:\s+background)?|"
    r"habilidades|compet[eê]ncias|skills|competencies|idiomas|languages|"
    r"certifica[cç][oõ]es|cursos|objetivo|resumo|summary|profile|perfil|"
    r"contato|contact|refer[eê]ncias|informa[cç][oõ]es\s+(?:adicionais|pessoais))\s*:?\s*$",
    re.IGNORECASE,
)

BULLET = re.compile(r"^(?:[-*•●▪►▸·∙‣]|\d{1,2}[\.\)\]])\s+")
YEAR = re.compile(r"\b(?:19|20)\d{2}\b")
SEPARATOR = re.compile(r"\s(?:\||—|–|/)\s|\s-\s")

ROLE_NOUN = re.compile(
    r"^(?:desenvolvedor(?:a)?|engenheir[oa]|analista|gerente|coordenador(?:a)?|"
    r"lider|estagiari[oa]|consultor(?:a)?|arquitet[oa]|designer|cientista|"
    r"especialista|supervisor(?:a)?|diretor(?:a)?|assistente|tecnico|"
    r"developer|engineer|analyst|manager|consultant|intern|"
    r"tech\s+lead|product\s+owner|scrum\s+master)\b"
)

ACTION_STEMS = (
    "desenvolv",
    "implement",
    "lider",
    "coorden",
    "gerenc",
    "geri ",
    "gerir",
    "criei",
    "criou",
    "criamos",
    "criar ",
    "criacao",
    "criando",
    "projetei",
    "projetar",
    "projetamos",
    "analisei",
    "analisamos",
    "analisar",
    "analisando",
    "analise ",
    "automat",
    "mantiv",
    "manuten",
    "mantendo",
    "otimiz",
    "constru",
    "particip",
    "elabor",
    "defini",
    "definicao",
    "acompanh",
    "entreg",
    "migrei",
    "migracao",
    "migrar",
    "migramos",
    "integrei",
    "integracao",
    "integrar",
    "integramos",
    "integrando",
    "apoiei",
    "apoio ",
    "treinei",
    "treinamento",
    "mentori",
    "document",
    "testei",
    "testes ",
    "publiquei",
    "publicacao",
    "configur",
    "administ",
    "monitor",
    "reduzi",
    "reducao",
    "aumentei",
    "aumento ",
    "implantei",
    "implantacao",
    "estrutur",
    "planej",
    "negoci",
    "atendi",
    "atendimento",
    "suporte ",
    "realizei",
    "realizacao",
    "executei",
    "execucao",
    "conduzi",
    "conduz",
    "organiz",
    "aprimor",
    "refator",
    "escrevi",
    "modelei",
    "modelagem",
    "arquitet",
    "responsavel",
    "colabor",
    "auxiliei",
    "auxiliar ",
    "auxilio ",
    "estabelec",
    "garanti",
    "assegur",
    "facilitei",
    "facilitacao",
    "mapeei",
    "mapeamento",
    "identifi",
    "corrigi",
    "correcao",
    "ajustei",
    "melhorei",
    "melhoria",
    "desenhei",
    "desenho ",
    "especifi",
    "homolog",
    "operei",
    "operacao ",
    "programei",
    "programacao",
    "atuacao",
    "atuei",
    "atuamos",
    "atuar ",
    "atuando",
    "atuou",
    "developed",
    "built ",
    "led ",
    "designed",
    "managed",
    "created",
    "improved",
    "maintained",
    "automated",
    "delivered",
    "migrated",
    "integrated",
    "owned ",
    "collaborated",
    "wrote ",
    "tested ",
    "deployed",
    "reduced",
    "increased",
    "architected",
    "mentored",
    "coordinated",
    "analyzed",
    "supported",
    "launched",
    "optimized",
    "refactored",
    "documented",
    "established",
    "facilitated",
    "shipped",
    "drove ",
)

MONTH = (
    r"jan(?:eiro)?|fev(?:ereiro)?|mar(?:ço|co)?|abr(?:il)?|mai(?:o)?|jun(?:ho)?|"
    r"jul(?:ho)?|ago(?:sto)?|set(?:embro)?|out(?:ubro)?|nov(?:embro)?|dez(?:embro)?|"
    r"feb|apr|may|aug|sep|oct"
)
DATE_TOKEN = re.compile(
    rf"\b(?:{MONTH}|(?:19|20)\d{{2}}|atual|presente|current|present|\d{{1,2}}/\d{{4}})\b",
    re.IGNORECASE,
)
DATE_FILLER = re.compile(
    r"\b(?:ate|até|a|to|de|do|da|e|o|em|desde|from)\b|[|—–/\-–.,()]",
    re.IGNORECASE,
)


class ResumeReadError(Exception):
    """O arquivo não pôde ser lido como currículo."""


def fold(text: str) -> str:
    normalized = unicodedata.normalize("NFD", text)
    without = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return without.casefold().strip()


def clean_space(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def read_resume(path: Path) -> str:
    suffix = path.suffix.casefold()
    data = path.read_bytes()
    if len(data) > MAX_BYTES:
        raise ResumeReadError("O arquivo passa de 8 MB.")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    raise ResumeReadError("Use um arquivo PDF ou DOCX.")


def read_resume_bytes(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.casefold()
    if len(data) > MAX_BYTES:
        raise ResumeReadError("O arquivo passa de 8 MB.")
    if not data:
        raise ResumeReadError("O arquivo está vazio.")
    if suffix == ".pdf":
        return _pdf_bytes(data)
    if suffix == ".docx":
        return _docx_bytes(data)
    raise ResumeReadError("Use um arquivo PDF ou DOCX.")


def list_activities(text: str) -> list[dict[str, str | list[str]]]:
    lines = _lines(text)
    sections = _sections(lines)
    groups: list[dict[str, str | list[str]]] = []
    for name, section_lines in sections:
        groups.extend(_parse_section(section_lines, name))
    if groups:
        return _dedupe(groups)
    fallback = _parse_section(
        [line for line in lines if not _is_header(line)],
        "Currículo",
    )
    return _dedupe(fallback)


def _lines(text: str) -> list[str]:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"[ \t]*[•●▪►▸][ \t]*", "\n• ", cleaned)
    return [clean_space(line) for line in cleaned.split("\n")]


def _is_header(line: str) -> bool:
    if SECTION_STOP.match(line):
        return True
    return any(pattern.match(line) for _, pattern in SECTION_STARTS)


def _sections(lines: list[str]) -> list[tuple[str, list[str]]]:
    current: str | None = None
    buckets: list[tuple[str, list[str]]] = []
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if current and buffer:
            buckets.append((current, buffer))
        buffer = []

    for line in lines:
        if not line:
            if current:
                buffer.append("")
            continue
        if SECTION_STOP.match(line):
            flush()
            current = None
            continue
        started = False
        for name, pattern in SECTION_STARTS:
            if pattern.match(line):
                flush()
                current = name
                started = True
                break
        if started:
            continue
        if current:
            buffer.append(line)
    flush()
    labels = {"experiencia": "Experiência", "projetos": "Projetos"}
    return [(labels[name], content) for name, content in buckets]


def _parse_section(lines: list[str], secao: str) -> list[dict[str, str | list[str]]]:
    groups: list[dict[str, str | list[str]]] = []
    current: dict[str, str | list[str]] | None = None
    pending_period = ""

    for index, raw in enumerate(lines):
        if not raw or _is_noise(raw):
            continue
        if _is_date_only(raw):
            period = _clean_period(raw)
            if current and not current["periodo"] and not current["atividades"]:
                current["periodo"] = period
            else:
                pending_period = period
            continue
        following = [line for line in lines[index + 1 : index + 5] if line]
        if _is_role(raw, following):
            context, period = _split_period(raw)
            if not period and pending_period:
                period = pending_period
            pending_period = ""
            current = {
                "secao": secao,
                "contexto": context,
                "periodo": period,
                "atividades": [],
            }
            groups.append(current)
            continue
        if _is_activity(raw):
            if current is None:
                current = {
                    "secao": secao,
                    "contexto": "",
                    "periodo": pending_period,
                    "atividades": [],
                }
                pending_period = ""
                groups.append(current)
            elif pending_period and not current["periodo"]:
                current["periodo"] = pending_period
                pending_period = ""
            activities = current["atividades"]
            assert isinstance(activities, list)
            activities.append(_clean_activity(raw))
            continue
        if current and current["atividades"] and _is_continuation(raw):
            activities = current["atividades"]
            assert isinstance(activities, list)
            activities[-1] = clean_space(f"{activities[-1]} {raw}")

    return [group for group in groups if group["atividades"]]


def _dedupe(groups: list[dict[str, str | list[str]]]) -> list[dict[str, str | list[str]]]:
    kept: list[dict[str, str | list[str]]] = []
    seen: set[tuple[str, str]] = set()
    for group in groups:
        activities = group["atividades"]
        assert isinstance(activities, list)
        unique: list[str] = []
        for activity in activities:
            key = (fold(str(group["contexto"])), fold(activity))
            if key in seen:
                continue
            seen.add(key)
            unique.append(activity)
        if unique:
            group["atividades"] = unique
            kept.append(group)
    return kept


def _is_noise(line: str) -> bool:
    if _is_date_only(line):
        return False
    folded = fold(line)
    if "@" in line or "http://" in folded or "https://" in folded:
        return True
    if "linkedin.com" in folded or "github.com" in folded:
        return True
    if _is_place(folded):
        return True
    compact = re.sub(r"\s", "", line)
    digits = re.sub(r"\D", "", line)
    if len(digits) >= 8 and compact and len(digits) / len(compact) >= 0.4:
        return True
    return folded in {
        "email",
        "e-mail",
        "telefone",
        "celular",
        "contato",
        "linkedin",
        "github",
        "curriculo",
        "remoto",
        "hibrido",
        "presencial",
        "home office",
    }


def _is_place(folded: str) -> bool:
    if re.fullmatch(r"[a-z\s]{3,28}(?:,|\s[—–\-]\s)[a-z]{2}", folded):
        return True
    return re.fullmatch(
        r"(?:remoto|hibrido|presencial|home office|brasil|brazil|sao paulo|"
        r"rio de janeiro|belo horizonte|curitiba|porto alegre|recife|salvador|"
        r"fortaleza|brasilia|lisboa|porto|remote|hybrid|on-?site)"
        r"(?:\s*[,—–\-/]\s*.*)?",
        folded,
    ) is not None


def _is_date_only(line: str) -> bool:
    if not DATE_TOKEN.search(line):
        return False
    remainder = DATE_TOKEN.sub(" ", line)
    remainder = DATE_FILLER.sub(" ", remainder)
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", remainder)
    return len(fold(letters)) <= 2


def _is_role(line: str, following: list[str]) -> bool:
    if _is_noise(line) or _is_date_only(line) or BULLET.match(line):
        return False
    if len(line) > 90 or line.endswith("."):
        return False
    folded = fold(line)
    has_separator = SEPARATOR.search(line) is not None
    if _starts_with_action(line) and not has_separator and not YEAR.search(line):
        return False
    if has_separator or ROLE_NOUN.match(folded):
        return True
    if YEAR.search(line):
        return True
    if len(line) <= 60 and _word_count(line) >= 3 and any(
        BULLET.match(item) or _is_date_only(item) for item in following[:3]
    ):
        return True
    return False


def _is_activity(line: str) -> bool:
    if _is_date_only(line) or _is_noise(line):
        return False
    content = _clean_activity(line)
    if len(content) < 12:
        return False
    if BULLET.match(line):
        return True
    return _starts_with_action(content)


def _is_continuation(line: str) -> bool:
    if _is_header(line) or BULLET.match(line) or _is_date_only(line) or _starts_with_action(line):
        return False
    return bool(line[:1].islower())


def _starts_with_action(line: str) -> bool:
    folded = fold(BULLET.sub("", line))
    if ROLE_NOUN.match(folded):
        return False
    return any(folded.startswith(stem) for stem in ACTION_STEMS)


def _split_period(line: str) -> tuple[str, str]:
    match = None
    for candidate in DATE_TOKEN.finditer(line):
        match = candidate
        break
    if not match:
        context = clean_space(SEPARATOR.sub(" — ", line).strip(" -—–|"))
        return context, ""
    context = clean_space(line[: match.start()])
    period = _clean_period(line[match.start() :])
    context = clean_space(SEPARATOR.sub(" — ", context).strip(" -—–|"))
    if len(fold(re.sub(r"[^A-Za-zÀ-ÿ]", "", context))) < 3:
        return clean_space(line), ""
    return context, period


def _clean_period(line: str) -> str:
    return clean_space(line.strip(" -—–|"))


def _clean_activity(line: str) -> str:
    text = BULLET.sub("", line).strip()
    text = clean_space(text).strip(" ;")
    if text and text[0].isalpha() and text[0].islower():
        text = text[0].upper() + text[1:]
    return text


def _word_count(line: str) -> int:
    return len(re.findall(r"[A-Za-zÀ-ÿ0-9]+", line))


def _read_pdf(path: Path) -> str:
    try:
        return _pdf_bytes(path.read_bytes())
    except ResumeReadError:
        raise
    except Exception as exc:  # noqa: BLE001 - superfície única para o leitor
        raise ResumeReadError("Não consegui ler o texto deste arquivo.") from exc


def _pdf_bytes(data: bytes) -> str:
    from io import BytesIO

    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ResumeReadError("O PDF está protegido por senha.")
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
    except ResumeReadError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ResumeReadError("Não consegui ler o texto deste arquivo.") from exc
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        return _docx_bytes(path.read_bytes())
    except ResumeReadError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ResumeReadError("Não consegui ler o texto deste arquivo.") from exc


def _docx_bytes(data: bytes) -> str:
    from io import BytesIO

    if not zipfile.is_zipfile(BytesIO(data)):
        raise ResumeReadError("Não consegui ler o texto deste arquivo.")
    try:
        document = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ResumeReadError("Não consegui ler o texto deste arquivo.") from exc
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = clean_space(paragraph.text)
        if not text:
            lines.append("")
            continue
        if _is_list_paragraph(paragraph):
            text = f"• {text}"
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = clean_space(cell.text)
                if value and value not in cells:
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
    if not any(line.strip() for line in lines):
        raise ResumeReadError("Não consegui ler o texto deste arquivo.")
    return "\n".join(lines)


def _is_list_paragraph(paragraph: object) -> bool:
    style = getattr(paragraph, "style", None)
    style_name = fold(getattr(style, "name", "") or "")
    if "list" in style_name:
        return True
    element = getattr(paragraph, "_p", None)
    properties = getattr(element, "pPr", None) if element is not None else None
    return properties is not None and properties.numPr is not None
